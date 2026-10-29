"""
Ekspor isi database BakeERP ke berkas yang mudah dibuka.

Menghasilkan:
  1. docs/Database_BakeERP.xlsx  -> semua tabel dalam satu workbook (1 sheet per tabel)
  2. docs/bakery_dump.sql        -> skema + isi tabel dalam format SQL (bisa dibuka/diimpor di DBeaver)

Jalankan:  python docs/export_db.py
"""

import os
import sqlite3
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import core

DOCS = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(DOCS, "Database_BakeERP.xlsx")
DUMP = os.path.join(DOCS, "bakery_dump.sql")

URUTAN = ["users", "supplier", "bahan", "produk", "resep", "pembelian", "pembelian_item",
          "produksi", "produksi_item", "penjualan", "penjualan_item", "pengeluaran", "kas", "audit"]


def info_tabel(conn):
    hasil = []
    for t in URUTAN:
        n = conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
        kolom = [r[1] for r in conn.execute(f"PRAGMA table_info({t})")]
        hasil.append({"tabel": t, "baris": n, "kolom": len(kolom), "nama_kolom": kolom})
    return hasil


def ke_excel(conn, info):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws["A1"] = "DATABASE BAKERP — Roti Kita Bakery"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Berkas: bakery.db (SQLite) · 14 tabel"
    ws["A2"].font = Font(italic=True, size=10)

    hdr = ["No", "Tabel", "Jumlah baris", "Jumlah kolom", "Kolom"]
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="5B3A1E")
        c.alignment = Alignment(horizontal="center")

    for i, t in enumerate(info, start=1):
        r = 4 + i
        ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=t["tabel"]).font = Font(bold=True)
        ws.cell(row=r, column=3, value=t["baris"]).alignment = Alignment(horizontal="right")
        ws.cell(row=r, column=4, value=t["kolom"]).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=5, value=", ".join(t["nama_kolom"]))
    for col, w in zip("ABCDE", (5, 20, 14, 14, 90)):
        ws.column_dimensions[col].width = w

    for t in URUTAN:
        cur = conn.execute(f"SELECT * FROM {t}")
        kolom = [d[0] for d in cur.description]
        sh = wb.create_sheet(t[:31])
        for i, h in enumerate(kolom, start=1):
            c = sh.cell(row=1, column=i, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="5B3A1E")
            c.alignment = Alignment(horizontal="center")
        for ri, row in enumerate(cur.fetchall(), start=2):
            for ci, val in enumerate(row, start=1):
                sh.cell(row=ri, column=ci, value=val)
        for i, h in enumerate(kolom, start=1):
            panjang = max(len(str(h)), 10)
            sh.column_dimensions[get_column_letter(i)].width = min(panjang + 4, 42)
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = f"A1:{get_column_letter(len(kolom))}{sh.max_row}"

    wb.save(XLSX)
    return XLSX


def ke_sql(conn):
    baris = ["-- Dump database BakeERP (SQLite)",
             f"-- Dibuat: {core.now_str()}",
             "-- Buka di DBeaver: File > Open File, atau import ke SQLite/MySQL",
             ""]
    for t in URUTAN:
        ddl = conn.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name=?", (t,)).fetchone()
        if ddl and ddl[0]:
            baris.append(f"-- ===== Tabel: {t} =====")
            baris.append(ddl[0].strip() + ";")
            baris.append("")
        cur = conn.execute(f"SELECT * FROM {t}")
        kolom = [d[0] for d in cur.description]
        rows = cur.fetchall()
        if rows:
            kol = ", ".join(kolom)
            for r in rows:
                nilai = []
                for v in r:
                    if v is None:
                        nilai.append("NULL")
                    elif isinstance(v, (int, float)):
                        nilai.append(str(v))
                    else:
                        nilai.append("'" + str(v).replace("'", "''") + "'")
                baris.append(f"INSERT INTO {t} ({kol}) VALUES ({', '.join(nilai)});")
            baris.append("")
    with open(DUMP, "w", encoding="utf-8") as f:
        f.write("\n".join(baris))
    return DUMP


def main():
    conn = sqlite3.connect(core.DB_PATH)
    conn.row_factory = sqlite3.Row
    info = info_tabel(conn)
    print("Database:", core.DB_PATH)
    print("-" * 56)
    total = 0
    for t in info:
        total += t["baris"]
        print(f"  {t['tabel']:<18} {t['baris']:>6} baris   {t['kolom']:>2} kolom")
    print("-" * 56)
    print(f"  {'TOTAL':<18} {total:>6} baris")
    print()
    print("Excel :", ke_excel(conn, info))
    print("SQL   :", ke_sql(conn))
    conn.close()


if __name__ == "__main__":
    main()
